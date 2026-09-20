# -*- coding: utf-8 -*-
"""众云信息科技电子产品回收综合服务平台V1.0 软著申请数据表"""
import sys, os

XLSX_SKILL_DIR = r'C:\Users\Administrator\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.4\skills\xlsx'
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, 'templates')]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import *  # noqa
from openpyxl import Workbook

OUT = r'D:\maintain\软著申报材料\回收综合服务平台V1.0\众云信息科技电子产品回收综合服务平台V1.0-申请数据.xlsx'

wb = Workbook()

# ==================== Sheet1 软著登记数据 ====================
ws = wb.active
ws.title = '软著登记数据'
fields = [
    ('软件全称', '众云信息科技电子产品回收综合服务平台'),
    ('软件简称', '电子产品回收平台（可不填）'),
    ('版本号', 'V1.0'),
    ('著作权人（申请人）', '深圳市众云信息科技有限公司'),
    ('统一社会信用代码', '【请填写】'),
    ('注册地址', '【请填写】'),
    ('联系人', '【请填写】'),
    ('联系电话', '【请填写】'),
    ('开发方式', '独立开发（企业）'),
    ('权利范围', '全部权利'),
    ('首次发表日期', '【请填写；未发表则填：未发表】'),
    ('开发完成日期', '2026年9月'),
    ('开发该软件的硬件环境', '8核CPU、16GB内存、512GB固态硬盘的PC兼容机'),
    ('软件运行的硬件环境', '服务端：4核CPU、8GB内存、40GB以上硬盘的服务器；用户端：支持微信小程序的智能手机、可运行现代浏览器的PC'),
    ('开发该软件的操作系统', 'Microsoft Windows 10 / 11'),
    ('软件开发环境 / 开发工具', 'Visual Studio Code、微信开发者工具、Node.js 20、Docker Desktop、MySQL Workbench'),
    ('该软件的运行平台 / 操作系统', 'Linux服务器（Docker容器化部署），经HTTPS统一网关对外服务；用户端为微信客户端（iOS / Android）与PC浏览器'),
    ('软件运行支撑环境 / 支持软件', 'Node.js 20运行时、Express框架、Vue3、Element Plus、ECharts、MySQL 8.0数据库、Nginx 1.24反向代理、Docker容器环境'),
    ('编程语言', 'JavaScript（服务端Node.js、前端Vue3、小程序端）、WXML / WXSS'),
    ('源程序量', '约10,100行（10100行，含服务端、前端与小程序回收模块）'),
    ('开发目的',
     '解决二手电子产品回收业务中报价口径不一、渠道分散、订单难追踪、估价不透明等痛点，'
     '建立覆盖配价管理、自助估价、AI辅助报价、订单处理、平台比价、数据统计与环保贡献量化的全流程线上化回收综合服务平台，'
     '提升回收企业运营效率与数字化管理水平。'),
    ('面向领域 / 行业', '电子产品回收、二手数码交易、循环经济与绿色环保回收行业'),
    ('软件的主要功能',
     '（1）小程序端：回收分类浏览、品牌型号选择、全局型号搜索（多关键词AND与无空格模糊匹配）、平台比价、'
     '七步问答式估价、AI大模型辅助估价（DeepSeek，失败自动降级本地估价）、回收地址选择、订单提交、内部人员免付款回收申请；'
     '（2）管理后台：数据看板（经营指标、趋势、分布、环保贡献量化）、数据大屏、回收订单管理（筛选、详情、报价、状态流转、手动建单、CSV导出）、'
     '设备配价库（分类/品牌/型号管理、单型号调价、批量调价、调价幅度保护与审计）、回收平台管理、采购链接管理、'
     '估价配置（二手基准系数、品类系数、六因子系数、估价模拟器）、操作日志、系统设置（三级角色权限、跨系统单点登录）。'),
    ('技术特点',
     '系统采用前后端分离与读写分离架构：管理后台（Node.js + Express + Vue3 + Element Plus）负责配价与配置写入，'
     '小程序服务端只读共享同一MySQL业务库并向小程序输出公开接口，两端独立部署、以库为界协作；'
     '配价目录小程序本地内置秒开，云端十分钟缓存增量更新并合并视觉资源，离线可用；'
     '估价采用多因子系数连乘模型（六因子×二手基准系数×品类系数），后台调参小程序免发版热更新；'
     'AI估价链路含提示词约束、JSON严格解析与三层降级；调价幅度超限服务端强制二次确认并记三层审计；'
     'JWT认证+三级角色中间件+跨系统SSO单点登录；全栈Docker容器化，经Nginx统一HTTPS网关子路径对外服务。'),
]
last_col = 4  # B序号 C字段 D内容
setup_sheet(ws, title='众云信息科技电子产品回收综合服务平台 软件著作权登记所需数据', last_col=last_col)
for col_idx, h in enumerate(['序号', '字段', '内容 / 说明'], start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)
for i, (k, v) in enumerate(fields):
    rn = 5 + i
    ws.cell(row=rn, column=2, value=i + 1)
    ws.cell(row=rn, column=3, value=k)
    ws.cell(row=rn, column=4, value=v)
    style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
auto_fit_columns(ws, min_width=8, max_width=80, header_row=4, data_start_row=5)
ws.column_dimensions['D'].width = 78
auto_fit_row_heights(ws, header_row=4, data_start_row=5)

# ==================== Sheet2 源程序统计 ====================
ws2 = wb.create_sheet('源程序统计')
rows2 = [
    ['管理后台服务端（Node.js + Express）', 'JavaScript', 12, 2367],
    ['管理后台前端（Vue3 + Element Plus）', 'JavaScript / Vue', 18, 4190],
    ['小程序回收模块（页面+样式+数据）', 'JavaScript / WXML / WXSS', 6, 3244],
    ['小程序服务端回收接口（含AI估价）', 'JavaScript', 1, 299],
]
last_col2 = 6
setup_sheet(ws2, title='源程序构成统计', last_col=last_col2)
for col_idx, h in enumerate(['模块', '主要技术', '文件数', '代码行数'], start=2):
    ws2.cell(row=4, column=col_idx, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=5)
for i, r in enumerate(rows2):
    rn = 5 + i
    for col_idx, v in enumerate(r, start=2):
        ws2.cell(row=rn, column=col_idx, value=v)
    style_data_row(ws2, row_num=rn, col_start=2, col_end=5, row_index=i)
tr = 5 + len(rows2)
ws2.cell(row=tr, column=2, value='合计')
ws2.cell(row=tr, column=4, value=sum(r[2] for r in rows2))
ws2.cell(row=tr, column=5, value=sum(r[3] for r in rows2))
style_total_row(ws2, row_num=tr, col_start=2, col_end=5)
for rn in range(5, tr + 1):
    ws2.cell(row=rn, column=4).number_format = FORMATS['integer']
    ws2.cell(row=rn, column=5).number_format = FORMATS['integer']
    ws2.cell(row=rn, column=4).alignment = align_number()
    ws2.cell(row=rn, column=5).alignment = align_number()
note = ws2.cell(row=tr + 2, column=2,
                value='注：行数为当前版本实际统计值；软著提交的源程序鉴别材料按规则取前30页+后30页（每页不少于50行）。')
note.font = font_caption()
auto_fit_columns(ws2, min_width=8, max_width=44, header_row=4, data_start_row=5)
auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

# ==================== Sheet3 申请材料清单 ====================
ws3 = wb.create_sheet('申请材料清单')
rows3 = [
    ['1', '软件著作权登记申请表', '在线填报（中国版权保护中心系统填写打印盖章）', '待办'],
    ['2', '著作权人身份证明', '企业营业执照副本复印件（加盖公章）', '待办'],
    ['3', '程序鉴别材料', '众云信息科技电子产品回收综合服务平台V1.0-源程序.pdf（60页，前30+后30，每页不少于50行）', '已完成'],
    ['4', '文档鉴别材料', '众云信息科技电子产品回收综合服务平台V1.0-说明书.docx（60页；插好截图后打印导出PDF提交，超60页时交前30+后30页）', '待补截图'],
    ['5', '说明书配图（可选）', '图表文件夹8张架构/流程图，可插入说明书对应章节', '已完成'],
    ['6', '经办人委托书', '如委托代理机构办理时需要', '按需'],
]
last_col3 = 6
setup_sheet(ws3, title='软件著作权登记申请材料清单', last_col=last_col3)
for col_idx, h in enumerate(['序号', '材料', '要求 / 说明', '状态'], start=2):
    ws3.cell(row=4, column=col_idx, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=5)
for i, r in enumerate(rows3):
    rn = 5 + i
    for col_idx, v in enumerate(r, start=2):
        ws3.cell(row=rn, column=col_idx, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=5, row_index=i)
auto_fit_columns(ws3, min_width=8, max_width=60, header_row=4, data_start_row=5)
ws3.column_dimensions['D'].width = 58
auto_fit_row_heights(ws3, header_row=4, data_start_row=5)

wb.properties.creator = 'Z.ai'
wb.save(OUT)
print('XLSX saved:', OUT)
